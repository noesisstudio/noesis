from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Guia-perfiles-sociales-Noesis.docx"

FOREST = "14463B"
TEAL = "2E8B74"
CREAM = "F4F1E8"
INK = "15211C"
MUTED = "5D6B66"
SAGE = "E4EFE9"
WHITE = "FFFFFF"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run(run, *, name="Calibri", size=11, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    table.autofit = False
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_shading(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    set_run(run, size=9, color=MUTED)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 30, FOREST, 0, 8),
        ("Subtitle", 14, MUTED, 0, 18),
        ("Heading 1", 16, FOREST, 18, 10),
        ("Heading 2", 13, TEAL, 14, 7),
        ("Heading 3", 12, FOREST, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run("NOESIS  |  PERFILES OFICIALES")
        set_run(run, size=8.5, color=TEAL, bold=True)

        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(0)
        lead = p.add_run("BYNOESIS.COM  ·  ")
        set_run(lead, size=8.5, color=MUTED, bold=True)
        add_page_field(p)


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run(run, size=30, color=FOREST, bold=True)
    if subtitle:
        p = doc.add_paragraph(style="Subtitle")
        run = p.add_run(subtitle)
        set_run(run, size=14, color=MUTED)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    size = {1: 16, 2: 13, 3: 12}[level]
    color = FOREST if level != 2 else TEAL
    set_run(run, size=size, color=color, bold=True)
    return p


def add_body(doc, text, *, bold_label=None, italic=False):
    p = doc.add_paragraph()
    if bold_label:
        label = p.add_run(bold_label)
        set_run(label, bold=True, color=FOREST)
    run = p.add_run(text)
    set_run(run, italic=italic)
    return p


def add_fact_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for index, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        set_cell_shading(cells[0], SAGE)
        if index % 2:
            set_cell_shading(cells[1], "FAF8F3")
        left = cells[0].paragraphs[0]
        right = cells[1].paragraphs[0]
        left.paragraph_format.space_after = Pt(0)
        right.paragraph_format.space_after = Pt(0)
        set_run(left.add_run(label), size=10, color=FOREST, bold=True)
        set_run(right.add_run(value), size=10.5, color=INK)
    set_table_geometry(table, [2700, 6660])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_copy_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    set_paragraph_shading(p, CREAM)
    for index, line in enumerate(text.split("\n")):
        run = p.add_run(line)
        set_run(run, name="Consolas", size=10, color=FOREST)
        if index < len(text.split("\n")) - 1:
            run.add_break()
    return p


def add_check_table(doc, checks):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    head = table.rows[0].cells
    for cell in head:
        set_cell_shading(cell, FOREST)
    set_run(head[0].paragraphs[0].add_run("Estado"), size=9.5, color=WHITE, bold=True)
    set_run(head[1].paragraphs[0].add_run("Comprobación"), size=9.5, color=WHITE, bold=True)
    for item in checks:
        cells = table.add_row().cells
        set_run(cells[0].paragraphs[0].add_run("Pendiente"), size=9.5, color=TEAL, bold=True)
        set_run(cells[1].paragraphs[0].add_run(item), size=9.5, color=INK)
    set_table_geometry(table, [1500, 7860])
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    set_run(p.add_run(text), size=8.5, color=MUTED, italic=True)


def add_page_break(doc):
    doc.add_page_break()


def build():
    doc = Document()
    configure_document(doc)

    # Cover: editorial guide with a restrained branded treatment.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(76)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT.parent / "logos" / "png" / "lockup" / "noesis-logo-primary-1024.png"), width=Inches(3.3))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(8)
    set_run(p.add_run("PERFILES OFICIALES"), size=10, color=TEAL, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("Instagram · Facebook · LinkedIn"), size=27, color=FOREST, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    set_run(p.add_run("Archivos, textos y configuración listos para publicar"), size=13, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    set_paragraph_shading(p, SAGE)
    set_run(p.add_run("Noesis lleva la oficina mientras tú haces el trabajo."), size=14, color=FOREST, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("Versión 1.1.2 · 31 de agosto de 2026"), size=9.5, color=MUTED)

    add_page_break(doc)
    add_title(doc, "Instagram", "Identificación, educación, humor de oficio y demostraciones breves")
    add_fact_table(doc, [
        ("Nombre", "Noesis | Gestión para autónomos"),
        ("Usuario", "@bynoesis"),
        ("Cuenta", "Empresa · pública"),
        ("Categoría", "Producto/servicio o Empresa de software"),
        ("Web", "https://bynoesis.com"),
        ("Correo", "info@bynoesis.com"),
    ])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / "instagram" / "SUBIR-imagen-perfil-instagram-1080.png"), width=Inches(1.55))
    add_caption(doc, "Archivo: SUBIR-imagen-perfil-instagram-1080.png · Instagram no utiliza banner")
    add_heading(doc, "Bio oficial", 2)
    add_copy_block(doc, "Tu negocio, en orden y bajo control.\nMenos papeleo. Más tiempo para tu oficio.\nPara autónomos de servicios.\n↓ Conoce Noesis")
    add_heading(doc, "Bio para el piloto", 2)
    add_copy_block(doc, "Noesis lleva la oficina mientras tú haces el trabajo.\nAgenda, facturas y documentos bajo control.\n↓ Solicita acceso")
    add_heading(doc, "Destacados iniciales", 2)
    add_body(doc, "Empieza · Cómo funciona · Autónomos · Facturas · Nosotros")

    add_page_break(doc)
    add_title(doc, "Facebook", "Confianza, comunidad, grupos profesionales y distribución")
    add_fact_table(doc, [
        ("Página", "Noesis"),
        ("Usuario", "@bynoesis"),
        ("Categoría", "Empresa de software · Servicio empresarial"),
        ("Web", "https://bynoesis.com"),
        ("Correo", "info@bynoesis.com"),
        ("Botón", "Más información mientras el acceso sea controlado"),
    ])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / "facebook" / "SUBIR-portada-facebook-1640x856.png"), width=Inches(6.15))
    add_caption(doc, "Portada: SUBIR-portada-facebook-1640x856.png · revisar recorte en móvil")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / "facebook" / "SUBIR-imagen-perfil-facebook-1080.png"), width=Inches(1.25))
    add_caption(doc, "Perfil: SUBIR-imagen-perfil-facebook-1080.png")
    add_heading(doc, "Presentación corta", 2)
    add_copy_block(doc, "Noesis lleva la oficina mientras tú haces el trabajo: agenda, clientes, facturas, documentos y control del negocio para autónomos de servicios.")
    add_heading(doc, "Descripción completa de Facebook", 1)
    add_copy_block(doc, "Noesis es la mano derecha del autónomo de servicios. Ayuda a mantener la agenda, los clientes, los trabajos, las facturas, los documentos, los costes y las tareas del negocio en orden, sin obligarte a aprender un ERP ni pasar el domingo haciendo papeleo.\n\nNoesis entiende qué está pasando, prepara lo administrativo y te enseña qué toca hacer. Tú mantienes el control y confirmas las acciones importantes.\n\nEstá pensado para autónomos y pequeños negocios de fontanería, electricidad, reformas, climatización, limpieza, mantenimiento, jardinería y otros servicios.\n\nMás información y acceso: https://bynoesis.com")
    add_heading(doc, "Configuración recomendada", 2)
    add_body(doc, "Mientras el acceso sea controlado, usar Más información. Cambiar a Registrarte solo cuando el alta pública esté validada y exista una URL real de registro.")
    add_check_table(doc, [
        "Crear una Página de empresa, no un perfil personal.",
        "Activar autenticación en dos pasos para los administradores.",
        "Añadir al menos dos administradores con cuentas distintas.",
        "Probar portada y botón en ordenador y móvil.",
    ])

    add_page_break(doc)
    add_title(doc, "LinkedIn", "Credibilidad empresarial, fundadores, producto y alianzas")
    add_fact_table(doc, [
        ("Página", "Noesis"),
        ("URL", "linkedin.com/company/bynoesis"),
        ("Sector", "Desarrollo de software"),
        ("Tamaño", "2-10 empleados, solo si refleja el equipo real"),
        ("Tipo", "Empresa privada, solo si refleja la forma jurídica real"),
        ("Web", "https://bynoesis.com"),
    ])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / "linkedin" / "SUBIR-portada-linkedin-4200x700.png"), width=Inches(6.2))
    add_caption(doc, "Portada: SUBIR-portada-linkedin-4200x700.png · especificación oficial de Página")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / "linkedin" / "SUBIR-logo-linkedin-400.png"), width=Inches(1.2))
    add_caption(doc, "Logotipo: SUBIR-logo-linkedin-400.png")
    add_heading(doc, "Lema", 2)
    add_copy_block(doc, "Noesis lleva la oficina mientras tú haces el trabajo.")

    add_page_break(doc)
    add_heading(doc, "Descripción de empresa de LinkedIn", 1)
    add_copy_block(doc, "Noesis es el copiloto operativo para autónomos y pequeños negocios de servicios. Conecta agenda, clientes, trabajos, proyectos, facturas, documentos, costes y gestoría para quitar ruido mental y devolver tiempo y control.\n\nNo es otro programa de facturación ni un chatbot genérico. Noesis entiende qué está pasando en el negocio, prepara el trabajo administrativo, explica lo importante y propone la siguiente acción. El profesional conserva el control y confirma las decisiones sensibles.\n\nNacemos para ayudar a fontaneros, electricistas, instaladores, reformas, climatización, limpieza, mantenimiento, jardinería y otros profesionales que tienen más trabajo que tiempo para gestionar la oficina.\n\nHaz tu trabajo; Noesis te ordena el negocio.")
    add_heading(doc, "Especialidades", 2)
    add_body(doc, "Gestión para autónomos · Software para empresas de servicios · Organización del negocio · Agenda y clientes · Facturación y documentos · Proyectos, costes y márgenes · Automatización con control humano · Experiencia web y WhatsApp")
    add_heading(doc, "Comprobación", 2)
    add_check_table(doc, [
        "Completar URL, sector, descripción, logotipo y portada.",
        "Añadir dos superadministradores y activar autenticación en dos pasos.",
        "Vincular la experiencia de los fundadores a la Página.",
        "Probar el botón Visitar sitio web y el recorte móvil.",
    ])

    add_page_break(doc)
    add_title(doc, "Canales adicionales", "Reservar la marca sin dispersar el lanzamiento")
    add_heading(doc, "YouTube", 2)
    add_body(doc, "Reservar @bynoesis y subir RESERVAR-avatar-youtube-800.png. Será el siguiente canal útil para demostraciones, tutoriales y entrevistas, pero no necesita calendario todavía.")
    add_heading(doc, "TikTok", 2)
    add_body(doc, "Reservar @bynoesis y subir RESERVAR-avatar-tiktok-1080.png. Activarlo solo cuando Instagram Reels tenga un formato repetible que pueda reutilizarse sin duplicar trabajo.")
    add_heading(doc, "No abrir todavía", 2)
    add_body(doc, "X/Twitter y Pinterest no concentran al cliente inicial. Google Business Profile solo debe crearse si Noesis cumple las condiciones de atención presencial o ubicación pública; nunca se inventa una dirección.")
    add_heading(doc, "Seguridad y propiedad", 1)
    add_check_table(doc, [
        "Registrar perfiles con correo corporativo controlado por la empresa.",
        "Activar autenticación en dos pasos y guardar códigos de recuperación.",
        "Añadir al menos dos administradores; ninguna cuenta depende de una persona.",
        "Guardar usuario, URL y propietario de cada perfil en el gestor de contraseñas.",
        "Enlazar los perfiles oficiales desde bynoesis.com cuando sus URL sean definitivas.",
    ])

    add_page_break(doc)
    add_title(doc, "Fuentes y especificaciones", "Fuentes oficiales consultadas")
    add_heading(doc, "Fuentes operativas", 2)
    add_body(doc, "LinkedIn - especificaciones de imágenes de Páginas: https://www.linkedin.com/help/linkedin/answer/a564330")
    add_body(doc, "LinkedIn - completar una Página: https://www.linkedin.com/help/linkedin/answer/a548448")
    add_body(doc, "Facebook - foto de perfil: https://www.facebook.com/help/163248423739693")
    add_body(doc, "Facebook - botón de acción: https://www.facebook.com/help/messenger-app/1638565256396310")
    add_body(doc, "Meta - Centro de cuentas: https://www.facebook.com/help/943858526073065")

    doc.core_properties.title = "Perfiles oficiales de Noesis"
    doc.core_properties.subject = "Instagram, Facebook y LinkedIn"
    doc.core_properties.author = "Noesis"
    doc.core_properties.keywords = "Noesis, branding, Instagram, Facebook, LinkedIn"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
